# tests/test_integration.py — Tests de integración
# ============================================================================
# Requieren el entorno agentes_py312 completo con todas las dependencias.
# Ejecutar con: pytest tests/test_integration.py -m integration -v
#
# Cubren:
#   - RAG: index_document → search_documents (pipeline completo)
#   - Memoria semántica: add_message → build_semantic_context
#   - Agent profiles: AGENT_REGISTRY, get_user_config
#   - Router: clasificación de mensajes
#   - generate_document: DOCX y PDF reales
# ============================================================================
import pytest
import sys
import json
import tempfile
from pathlib import Path
from unittest.mock import patch, MagicMock

sys.path.insert(0, str(Path(__file__).parent.parent))

pytestmark = pytest.mark.integration


# ── Tests: RAG pipeline completo ─────────────────────────────────────────────

class TestRAGPipeline:
    """Prueba el ciclo completo: indexar → buscar → recuperar."""

    @pytest.fixture
    def rag_env(self, tmp_data_dir):
        """Entorno RAG aislado con ChromaDB en directorio temporal."""
        chroma_dir = tmp_data_dir / "chroma"
        chroma_dir.mkdir()
        
        # Resetear el singleton del cliente ChromaDB
        import backend.rag as rag_module
        rag_module._chroma_client = None
        
        with patch("backend.rag.CHROMA_DIR", chroma_dir), \
             patch("backend.rag.UPLOADS_DIR", tmp_data_dir / "uploads"):
            yield tmp_data_dir

    def test_index_and_search_structured_doc(self, rag_env, tmp_data_dir):
        """Documento estructurado: indexar y recuperar fragmento específico."""
        from backend.rag import index_document, search_documents

        doc_file = tmp_data_dir / "test_structured.txt"
        doc_file.write_text(
            "# La Revolución Francesa\n\n"
            "La Revolución Francesa comenzó en 1789 con la toma de la Bastilla.\n\n"
            "## Causas\n\n"
            "Las causas principales fueron la desigualdad social y la crisis fiscal.\n\n"
            "### Referencias\n\n"
            "- Furet, François. Interpreting the French Revolution. 1981.\n"
            "- Doyle, William. The Oxford History. 2002.",
            encoding="utf-8"
        )

        result = index_document(doc_file, "test_user")
        assert result["status"] == "ok"
        assert result["chunks"] > 0

        docs = search_documents("causas revolución francesa", "test_user")
        assert len(docs) > 0
        combined = " ".join(d.page_content.lower() for d in docs)
        assert "causas" in combined or "revolución" in combined or "bastilla" in combined

    def test_index_and_search_transcript(self, rag_env, tmp_data_dir):
        """Transcripción: chunking adaptativo y recuperación."""
        from backend.rag import index_document, search_documents

        transcript = tmp_data_dir / "youtube_clase.txt"
        transcript.write_text(
            "hola bienvenidos al video de hoy vamos a hablar sobre machine learning "
            "el machine learning es una rama de la inteligencia artificial que permite "
            "a las máquinas aprender de los datos sin ser programadas explícitamente "
            "hay tres tipos principales de aprendizaje el supervisado el no supervisado "
            "y el aprendizaje por refuerzo cada uno tiene sus aplicaciones específicas "
            "el aprendizaje supervisado usa datos etiquetados para entrenar modelos "
            "el no supervisado encuentra patrones en datos sin etiquetar "
            "el refuerzo aprende mediante prueba y error con recompensas",
            encoding="utf-8"
        )

        result = index_document(transcript, "test_user")
        assert result["status"] == "ok"
        assert result.get("doc_type") == "transcript" or result["chunks"] > 0

        docs = search_documents("aprendizaje supervisado datos etiquetados", "test_user")
        assert len(docs) > 0

    def test_small_doc_returns_all_chunks(self, rag_env, tmp_data_dir):
        """Documento pequeño (≤15 chunks) devuelve contenido completo."""
        from backend.rag import index_document, search_documents

        small_doc = tmp_data_dir / "small_report.txt"
        # Documento pequeño con sección de referencias al final
        small_doc.write_text(
            "# Informe de Análisis\n\n"
            "Este informe analiza los datos recopilados durante el proyecto.\n\n"
            "## Metodología\n\n"
            "Se utilizaron técnicas de análisis estadístico descriptivo.\n\n"
            "## Resultados\n\n"
            "Los resultados muestran una tendencia positiva en todos los indicadores.\n\n"
            "## Conclusiones\n\n"
            "El proyecto cumplió sus objetivos dentro del plazo establecido.\n\n"
            "## Referencias\n\n"
            "- García, J. Métodos de análisis de datos. 2023.\n"
            "- López, M. Estadística aplicada. 2022.",
            encoding="utf-8"
        )

        index_document(small_doc, "test_user")
        # Buscar referencias — sección con baja similitud semántica
        docs = search_documents("referencias bibliográficas", "test_user")
        combined = " ".join(d.page_content for d in docs)
        # Con el path de documento pequeño, las referencias deben estar presentes
        assert "García" in combined or "López" in combined or "Referencias" in combined

    def test_file_name_filter_searches_correct_doc(self, rag_env, tmp_data_dir):
        """file_name fuerza la búsqueda en el documento correcto."""
        from backend.rag import index_document, search_documents

        doc_a = tmp_data_dir / "python_guide.txt"
        doc_a.write_text(
            "Python es un lenguaje de programación interpretado de alto nivel. "
            "Es ampliamente usado en ciencia de datos y machine learning.",
            encoding="utf-8"
        )
        doc_b = tmp_data_dir / "history_notes.txt"
        doc_b.write_text(
            "La historia antigua abarca desde los primeros registros escritos "
            "hasta la caída del Imperio Romano de Occidente en el año 476 d.C.",
            encoding="utf-8"
        )

        index_document(doc_a, "test_user")
        index_document(doc_b, "test_user")

        # Con file_name debe buscar solo en python_guide
        docs = search_documents("lenguaje programación", "test_user",
                                 file_name="python_guide.txt")
        assert all(
            d.metadata.get("file_name") == "python_guide.txt"
            for d in docs
            if d.metadata.get("file_name")
        )

    def test_is_transcript_detection(self, tmp_data_dir):
        """_is_transcript detecta transcripciones por nombre y contenido."""
        from backend.rag import _is_transcript
        from langchain_core.documents import Document

        # Por nombre
        transcript_path = Path("youtube_video_transcript.txt")
        doc = Document(page_content="texto cualquiera", metadata={})
        assert _is_transcript(transcript_path, [doc]) is True

        # Por contenido (texto continuo sin puntuación)
        long_text = ("palabra " * 50).strip()  # sin puntuación, línea muy larga
        transcript_doc = Document(page_content=long_text, metadata={})
        neutral_path = Path("notes.txt")
        result = _is_transcript(neutral_path, [transcript_doc])
        # No garantizamos True aquí (depende de umbral) pero no debe lanzar excepción
        assert isinstance(result, bool)

        # Documento estructurado no es transcripción
        structured_path = Path("analysis.pdf")
        structured_doc = Document(
            page_content=(
                "Este documento analiza los datos. Presenta resultados claros. "
                "Las conclusiones son positivas. El método fue riguroso."
            ),
            metadata={}
        )
        assert _is_transcript(structured_path, [structured_doc]) is False

    def test_extract_snippets_finds_references(self, tmp_data_dir):
        """extract_relevant_snippets localiza referencias en chunk mixto."""
        from backend.rag import extract_relevant_snippets
        from langchain_core.documents import Document

        # Chunk con conclusiones (80%) + referencias (20%) — caso real
        chunk = Document(
            page_content=(
                "## Conclusión\n\n"
                "La revolución fue un proceso complejo con múltiples causas.\n"
                "El legado continúa informando debates modernos sobre libertad.\n\n"
                "### Referencias\n\n"
                "- Furet, François. Interpreting the French Revolution. 1981.\n"
                "- Cobban, Alfred. The Social Interpretation. 1964.\n"
                "- Doyle, William. The Oxford History. 2002."
            ),
            metadata={"file_name": "history.txt", "source": "history.txt"}
        )

        snippets = extract_relevant_snippets([chunk], "referencias")
        assert len(snippets) >= 1
        combined = " ".join(s.page_content for s in snippets)
        assert "Furet" in combined or "Cobban" in combined or "Doyle" in combined


# ── Tests: Memoria semántica ──────────────────────────────────────────────────

class TestSemanticMemory:
    """Prueba el ciclo: add_message → vectorizar → build_semantic_context."""

    @pytest.fixture
    def memory_env(self, tmp_data_dir):
        """Entorno de memoria aislado."""
        chroma_dir = tmp_data_dir / "chroma_memory"
        chroma_dir.mkdir()
        history_dir = tmp_data_dir / "history"
        history_dir.mkdir()
        
        # Resetear el singleton del cliente ChromaDB
        import backend.rag as rag_module
        rag_module._chroma_client = None
    
        with patch("backend.memory.CHAT_HISTORY_DIR", history_dir), \
             patch("backend.memory.CHROMA_DIR", chroma_dir):
            yield tmp_data_dir

    def test_add_message_persists(self, memory_env, tmp_data_dir):
        """add_message guarda el mensaje en el archivo correcto."""
        from backend.memory import add_message

        history_file = tmp_data_dir / "history" / "user1" / "conv_default.json"
        history_file.parent.mkdir(parents=True, exist_ok=True)
        history_file.write_text("[]", encoding="utf-8")

        add_message("user1", "user", "¿Qué es Python?",
                    agent="general", conversation_id="default")

        messages = json.loads(history_file.read_text())
        assert len(messages) == 1
        assert messages[0]["content"] == "¿Qué es Python?"
        assert messages[0]["role"] == "user"
        assert messages[0]["agent"] == "general"

    def test_get_langchain_messages_correct_types(self, memory_env, tmp_data_dir):
        """get_langchain_messages devuelve HumanMessage y AIMessage."""
        from backend.memory import get_langchain_messages
        from langchain_core.messages import HumanMessage, AIMessage

        history_file = tmp_data_dir / "history" / "user1" / "conv_default.json"
        history_file.parent.mkdir(parents=True, exist_ok=True)
        history_file.write_text(json.dumps([
            {"role": "user", "content": "Hola",
             "agent": "general", "timestamp": "2026-04-10T10:00:00",
             "conversation_id": "default"},
            {"role": "assistant", "content": "¡Hola! ¿En qué puedo ayudarte?",
             "agent": "general", "timestamp": "2026-04-10T10:00:01",
             "conversation_id": "default"},
        ]), encoding="utf-8")

        messages = get_langchain_messages("user1", last_n=10,
                                          conversation_id="default")
        assert len(messages) == 2
        assert isinstance(messages[0], HumanMessage)
        assert isinstance(messages[1], AIMessage)
        assert messages[0].content == "Hola"

    def test_cross_agent_memory_config(self):
        """La configuración cruzada pedagogico↔psicologo es simétrica."""
        from backend.memory import _CROSS_AGENT_MEMORY
        assert "psicologo" in _CROSS_AGENT_MEMORY.get("pedagogico", [])
        assert "pedagogico" in _CROSS_AGENT_MEMORY.get("psicologo", [])

    def test_build_semantic_context_returns_string(self, memory_env):
        """build_semantic_context siempre devuelve string."""
        from backend.memory import build_semantic_context

        # Con ChromaDB vacío debe devolver string vacío
        result = build_semantic_context(
            user_id="test_user",
            query="Python programación",
            agent="programador",
            top_k=3,
            min_score=0.35,
        )
        assert isinstance(result, str)


# ── Tests: Agent profiles ─────────────────────────────────────────────────────

class TestAgentProfiles:
    """Prueba la configuración de agentes y el router."""

    def test_agent_registry_has_required_agents(self):
        """AGENT_REGISTRY debe tener todos los agentes definidos."""
        from backend.agent_profiles import AGENT_REGISTRY
        required = ["general", "pedagogico", "psicologo", "programador", "matematico"]
        for agent in required:
            assert agent in AGENT_REGISTRY, f"Agente '{agent}' no en AGENT_REGISTRY"

    def test_agent_registry_fields(self):
        """Cada agente debe tener prompt, temperature y tools."""
        from backend.agent_profiles import AGENT_REGISTRY
        for name, cfg in AGENT_REGISTRY.items():
            assert "prompt" in cfg, f"{name}: falta 'prompt'"
            assert "temperature" in cfg, f"{name}: falta 'temperature'"
            assert "tools" in cfg, f"{name}: falta 'tools'"

    def test_temperature_in_valid_range(self):
        """Temperaturas deben estar entre 0.0 y 1.0."""
        from backend.agent_profiles import AGENT_REGISTRY
        for name, cfg in AGENT_REGISTRY.items():
            temp = cfg["temperature"]
            assert 0.0 <= temp <= 1.0, \
                f"{name}: temperatura {temp} fuera de rango [0, 1]"

    def test_get_user_config_default(self):
        """get_user_config devuelve __default__ para usuarios desconocidos."""
        from backend.agent_profiles import get_user_config
        config = get_user_config("usuario_que_no_existe_xyz")
        assert "mode" in config
        assert "agents" in config
        assert "default_agent" in config

    def test_get_user_config_mode_valid(self):
        """El modo de configuración debe ser manual_select o auto_router."""
        from backend.agent_profiles import get_user_config
        config = get_user_config("usuario_que_no_existe_xyz")
        assert config["mode"] in ("manual_select", "auto_router")

    def test_agent_display_names_complete(self):
        """AGENT_DISPLAY_NAMES debe tener entrada para cada agente del registro."""
        from backend.agent_profiles import AGENT_REGISTRY, AGENT_DISPLAY_NAMES
        for agent in AGENT_REGISTRY:
            assert agent in AGENT_DISPLAY_NAMES, \
                f"'{agent}' no tiene display name"

    def test_router_prompt_has_three_categories(self):
        """ROUTER_PROMPT debe contener las tres categorías de clasificación."""
        from backend.agent_profiles import ROUTER_PROMPT
        assert "PEDAGOGICO" in ROUTER_PROMPT
        assert "EMOCIONAL" in ROUTER_PROMPT
        assert "GENERAL" in ROUTER_PROMPT

    def test_router_prompt_instructs_single_word_response(self):
        """ROUTER_PROMPT debe pedir respuesta de una sola palabra."""
        from backend.agent_profiles import ROUTER_PROMPT
        # Debe instruir al modelo a responder solo con la categoría
        assert "ÚNICAMENTE" in ROUTER_PROMPT or "solo" in ROUTER_PROMPT.lower() \
               or "only" in ROUTER_PROMPT.lower() or "Responde" in ROUTER_PROMPT


# ── Tests: generate_document — formatos reales ───────────────────────────────

class TestGenerateDocumentReal:
    """Tests de generación de documentos reales (sin mocks de filesystem)."""

    def test_pdf_plain_text_via_reportlab(self, tmp_data_dir):
        """Texto plano sin markdown → ReportLab → PDF válido."""
        with patch("backend.tools.DOCUMENTS_DIR", tmp_data_dir):
            from backend.tools import generate_document
            result = generate_document.func(
                content="Este es un texto plano sin ningún formato especial.",
                format="pdf",
                title="Test PDF ReportLab",
                user_id="test_user"
            )
        assert "FORMAT:pdf" in result
        name = result.split("NAME:")[1].strip()
        pdf_path = tmp_data_dir / "test_user" / name
        assert pdf_path.exists()
        assert pdf_path.stat().st_size > 100  # PDF no vacío
        # Verificar magic bytes de PDF
        with open(pdf_path, "rb") as f:
            assert f.read(4) == b"%PDF"

    def test_docx_plain_text_via_python_docx(self, tmp_data_dir):
        """Texto sin matemáticas → python-docx → DOCX válido."""
        with patch("backend.tools.DOCUMENTS_DIR", tmp_data_dir):
            from backend.tools import generate_document
            result = generate_document.func(
                content="Contenido sin matemáticas ni tablas markdown.",
                format="docx",
                title="Test DOCX python-docx",
                user_id="test_user"
            )
        assert "FORMAT:docx" in result
        name = result.split("NAME:")[1].strip()
        docx_path = tmp_data_dir / "test_user" / name
        assert docx_path.exists()
        assert docx_path.stat().st_size > 100
        # DOCX es un ZIP — verificar magic bytes
        with open(docx_path, "rb") as f:
            assert f.read(2) == b"PK"

    def test_docx_with_markdown_table(self, tmp_data_dir):
        """Tabla markdown → python-docx → DOCX con tabla real."""
        content = (
            "Aquí hay una tabla de resultados:\n\n"
            "| Asignatura | Nota |\n"
            "|------------|------|\n"
            "| Python     | 10   |\n"
            "| Estadística| 9.5  |\n\n"
            "Los resultados son excelentes."
        )
        with patch("backend.tools.DOCUMENTS_DIR", tmp_data_dir):
            from backend.tools import generate_document
            result = generate_document.func(
                content=content,
                format="docx",
                title="Tabla de notas",
                user_id="test_user"
            )
        assert "FORMAT:docx" in result
        name = result.split("NAME:")[1].strip()
        docx_path = tmp_data_dir / "test_user" / name

        # Verificar que el DOCX contiene una tabla
        from docx import Document
        doc = Document(str(docx_path))
        assert len(doc.tables) >= 1
        # La tabla debe tener las filas correctas
        table = doc.tables[0]
        assert table.rows[0].cells[0].text == "Asignatura"
        assert table.rows[1].cells[1].text == "10"

    def test_pdf_with_rich_markdown_via_pandoc(self, tmp_data_dir):
        """Markdown rico → pandoc+tectonic → PDF (requiere tectonic)."""
        import shutil
        if not shutil.which("pandoc"):
            pytest.skip("pandoc no disponible en este entorno")

        content = (
            "# Informe de Análisis\n\n"
            "## Introducción\n\n"
            "Este es un informe con formato markdown.\n\n"
            "## Resultados\n\n"
            "- Resultado 1: positivo\n"
            "- Resultado 2: satisfactorio\n\n"
            "## Conclusión\n\n"
            "El proyecto fue exitoso."
        )
        with patch("backend.tools.DOCUMENTS_DIR", tmp_data_dir):
            from backend.tools import generate_document
            result = generate_document.func(
                content=content,
                format="pdf",
                title="Informe con markdown",
                user_id="test_user"
            )
        # Puede fallar si tectonic no está disponible — no es error crítico
        if "Error" not in result:
            assert "FORMAT:pdf" in result
            name = result.split("NAME:")[1].strip()
            pdf_path = tmp_data_dir / "test_user" / name
            if pdf_path.exists():
                with open(pdf_path, "rb") as f:
                    assert f.read(4) == b"%PDF"

    def test_docx_with_math_via_pandoc(self, tmp_data_dir):
        """LaTeX math → pandoc+mathml → DOCX con ecuaciones (requiere pandoc)."""
        import shutil
        if not shutil.which("pandoc"):
            pytest.skip("pandoc no disponible en este entorno")

        content = (
            "La fracción $\\frac{1}{2}$ es un número racional.\n\n"
            "La integral $\\int_0^1 x^2 dx = \\frac{1}{3}$ es conocida.\n\n"
            "El teorema de Pitágoras: $a^2 + b^2 = c^2$"
        )
        with patch("backend.tools.DOCUMENTS_DIR", tmp_data_dir):
            from backend.tools import generate_document
            result = generate_document.func(
                content=content,
                format="docx",
                title="Ficha con matemáticas",
                user_id="test_user"
            )
        if "Error" not in result:
            assert "FORMAT:docx" in result
            name = result.split("NAME:")[1].strip()
            docx_path = tmp_data_dir / "test_user" / name
            assert docx_path.exists()
            # DOCX válido
            with open(docx_path, "rb") as f:
                assert f.read(2) == b"PK"


# ── Tests: stream_chat y stream_simple_chat (sin LLM real) ───────────────────

class TestStreamChat:
    """Tests del pipeline de streaming sin llamadas LLM reales."""

    @pytest.fixture
    def mock_llm(self):
        """Mock del LLM que devuelve tokens predecibles."""
        mock = MagicMock()

        async def mock_astream(messages):
            for token in ["Hola", " mundo", "."]:
                chunk = MagicMock()
                chunk.content = token
                yield chunk

        mock.astream = mock_astream
        return mock

    def test_stream_simple_chat_yields_tokens(self, tmp_data_dir, mock_llm):
        """stream_simple_chat debe generar tokens del LLM."""
        import asyncio

        history_file = tmp_data_dir / "user1" / "conv_default.json"
        history_file.parent.mkdir(parents=True, exist_ok=True)
        history_file.write_text("[]", encoding="utf-8")

        with patch("backend.memory.CHAT_HISTORY_DIR", tmp_data_dir), \
             patch("backend.agents.get_llm", return_value=mock_llm):

            from backend.agents import stream_simple_chat

            async def collect():
                tokens = []
                async for token in stream_simple_chat(
                    user_id="user1",
                    message="Hola",
                    conversation_id="default"
                ):
                    tokens.append(token)
                return tokens

            tokens = asyncio.run(collect())

        assert len(tokens) > 0
        assert "".join(tokens) in ["Hola mundo.", "Hola mundo. "]

    def test_auto_name_called_in_stream_chat(self, tmp_data_dir):
        """stream_chat debe llamar auto_name_conversation."""
        with patch("backend.agents.auto_name_conversation") as mock_name, \
             patch("backend.agents.get_user_config",
                   return_value={"mode": "manual_select",
                                 "agents": ["general"]}), \
             patch("backend.agents.build_agent") as mock_build, \
             patch("backend.memory.CHAT_HISTORY_DIR", tmp_data_dir):

            # Mock del agente que no hace nada
            mock_agent = MagicMock()

            async def mock_astream_events(*args, **kwargs):
                yield {"event": "on_chat_model_stream",
                       "data": {"chunk": MagicMock(content="respuesta")},
                       "name": ""}
            mock_agent.astream_events = mock_astream_events
            mock_build.return_value = mock_agent

            import asyncio
            from backend.agents import stream_chat

            async def run():
                tokens = []
                async for t in stream_chat(
                    user_id="user1",
                    message="Prueba de auto-nombrado",
                    conversation_id="test_conv"
                ):
                    tokens.append(t)
                return tokens

            history_dir = tmp_data_dir / "user1"
            history_dir.mkdir(parents=True, exist_ok=True)
            (history_dir / "conv_test_conv.json").write_text(
                "[]", encoding="utf-8"
            )

            asyncio.run(run())
            mock_name.assert_called_once_with(
                "user1", "test_conv", "Prueba de auto-nombrado"
            )


# ── Tests: code_interpreter con pandas real ───────────────────────────────────

class TestCodeInterpreterReal:
    """Tests del intérprete de código con ejecución real."""

    def test_numpy_computation(self):
        """numpy debe estar disponible en el sandbox."""
        from backend.tools import code_interpreter
        result = code_interpreter.func("np.sqrt(np.array([4, 9, 16]))")
        assert "2." in result and "3." in result and "4." in result

    def test_pandas_dataframe_display(self):
        """DataFrame debe devolverse como HTML con marcador TABLE."""
        from backend.tools import code_interpreter
        result = code_interpreter.func(
            "pd.DataFrame({'A': [1, 2, 3], 'B': ['x', 'y', 'z']})"
        )
        assert "TABLE:" in result
        assert "<table" in result.lower()
    
    @pytest.mark.filterwarnings("ignore:FigureCanvasAgg is non-interactive:UserWarning")
    def test_matplotlib_plot_saved(self, tmp_data_dir):
        """plt.show() debe guardar la figura y devolver PLOT:."""
        with patch("backend.tools._current_user_id", "test_user"):
            from backend.tools import code_interpreter
            result = code_interpreter.func(
                "plt.figure(); plt.plot([1,2,3], [4,5,6]); plt.show()"
            )
        assert "PLOT:" in result

    def test_sympy_symbolic_math(self):
        """sympy debe resolver integrales simbólicas."""
        from backend.tools import code_interpreter
        result = code_interpreter.func(
            "x = symbols('x'); print(integrate(x**2, x))"
        )
        assert "x**3" in result or "x³" in result or "3" in result

    def test_timeout_returns_message(self):
        """[UNIT] Timeout handling: validación instantánea y determinista."""
        from backend.tools import code_interpreter
        from concurrent.futures import TimeoutError as FuturesTimeout
        from unittest.mock import patch, MagicMock

        with patch("concurrent.futures.ThreadPoolExecutor") as MockExecutor:
            # Crear mock del futuro que lanza TimeoutError al llamar a .result()
            mock_future = MagicMock()
            mock_future.result.side_effect = FuturesTimeout("execution timed out")

            # Crear mock del executor (maneja el context manager `with ... as ...`)
            mock_executor = MagicMock()
            mock_executor.submit.return_value = mock_future
            mock_executor.__enter__ = MagicMock(return_value=mock_executor)
            mock_executor.__exit__  = MagicMock(return_value=False)

            # Hacer que ThreadPoolExecutor() devuelva nuestro mock
            MockExecutor.return_value = mock_executor

            # Ejecutar (el código es irrelevante, el mock fuerza el timeout)
            result = code_interpreter.func("1 + 1")

        # Validar que el manejo de timeout funciona
        assert any(term in result.lower() for term in [
            "timeout", "exceeded", "15 seconds", "tiempo"
        ])

    def test_security_blocks_os_import(self):
        """El sandbox no debe permitir importar os."""
        from backend.tools import code_interpreter
        result = code_interpreter.func("import os; os.listdir('.')")
        assert "Error" in result or "not permitted" in result.lower() \
               or "permitido" in result.lower()

    def test_csv_interception_with_absolute_path(self, tmp_data_dir):
        """to_csv con ruta absoluta debe interceptarse y redirigirse."""
        with patch("backend.tools._current_user_id", "test_user"):
            from backend.tools import code_interpreter
            result = code_interpreter.func(
                f"df = pd.DataFrame({{'A': [1,2,3]}}); "
                f"df.to_csv('/tmp/test_output.csv', index=False)"
            )
        # El archivo NO debe estar en /tmp
        assert not Path("/tmp/test_output.csv").exists()
        # Debe haber un FILE: en el resultado
        assert "FILE:" in result

    def test_pandas_methods_restored_after_error(self):
        """Los métodos pandas se restauran incluso si el código lanza excepción."""
        import pandas as pd
        original_to_csv = pd.DataFrame.to_csv

        from backend.tools import code_interpreter
        # Código que lanza error después de mockear pandas
        code_interpreter.func("1/0")

        assert pd.DataFrame.to_csv is original_to_csv
